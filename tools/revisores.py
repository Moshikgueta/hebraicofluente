#!/usr/bin/env python3
"""revisores.py - as planilhas dos revisores nativos, de ida e de volta.

Etapa 4 do PLANO-INGLES.md. O texto em inglês do site sai em quatro
planilhas (.xlsx), uma linha por texto, com a coluna amarela "Your version"
para o revisor escrever. A volta lê as planilhas preenchidas e põe cada
versão no lugar de onde a linha saiu:

  tela:<arquivo>:<chave>   app/src/i18n/telas/<arquivo>.ts(x), no bloco `en`
  <arquivo>:<caminho>      data/en/<arquivo>.json, o campo `en` do registro
  <arquivo>:/<ponteiro>    transliteração: nunca aplicada sozinha, porque a do
                           inglês sai de uma regra (tools/translit-en.mjs) e
                           uma correção pede mudar a regra ou abrir exceção

Uma versão só entra quando:
  - o "English now" da planilha ainda é o texto de hoje (senão o inglês mudou
    depois de a planilha sair, e quem decide é uma pessoa);
  - o hebraico e as marcas {{...}} são os mesmos, na mesma ordem;
  - as variáveis {n}, {nome}... são as mesmas;
  - na tela, o valor de hoje é uma string simples numa linha só (função,
    JSX e texto em várias linhas vão para a lista do que aplicar à mão).

Tudo o que não entra, e todo comentário ou "Question" dos revisores, vai
para o relatório em Markdown.

Uso (da raiz):
  python3 tools/revisores.py gerar <pasta>               as quatro planilhas
  python3 tools/revisores.py aplicar <x.xlsx>... [--gravar] [--relatorio r.md]
  python3 tools/revisores.py ensaiar                      prova a ida e a volta

Sem --gravar, `aplicar` só mostra o que faria. Precisa do openpyxl e do Node.
"""

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

RAIZ = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
HEBRAICO = re.compile(r'\{\{[^}]*\}\}|[֐-׿]+')
VARIAVEL = re.compile(r'(?<!\{)\{[A-Za-z_][A-Za-z0-9_]*\}(?!\})')


def dados(raiz):
    """O texto de hoje, pelo tools/revisores-dados.mjs."""
    with tempfile.NamedTemporaryFile(suffix='.json', delete=False) as f:
        caminho = f.name
    try:
        subprocess.run(['node', os.path.join(raiz, 'tools', 'revisores-dados.mjs'), caminho],
                       check=True, cwd=raiz, stdout=subprocess.DEVNULL)
        with open(caminho, encoding='utf-8') as f:
            return json.load(f)
    finally:
        os.unlink(caminho)


# ── ida: as planilhas ─────────────────────────────────────────────────────

LEIA = [
    ("How to review", True),
    ("Thank you for helping! This file holds the English text of the Hebraico Fluente course, "
     "a course that teaches English speakers to read and write Hebrew.", False),
    ("", False),
    ("1. Read the column \"English now\". That is exactly what students will see.", False),
    ("2. If it is fine, leave the row alone (or pick OK in the \"Status\" column).", False),
    ("3. If it should change, write your version in the yellow column \"Your version\". "
     "Write the whole sentence, not just the changed word.", False),
    ("4. Use \"Comment\" for anything else: a doubt, a term that sounds odd, an explanation "
     "that does not make sense.", False),
    ("", False),
    ("Please keep exactly as they are", True),
    ("- Hebrew text, and anything between double braces, like {{א}}. The site shows Hebrew there.", False),
    ("- Words between single braces, like {n} or {name}. The site fills them in (a number, a name, a link).", False),
    ("- The symbol ◆ means \"XP points\".", False),
    ("", False),
    ("Good to know", True),
    ("- The column \"Portuguese (original)\" is the source text. You do not need to read it; "
     "it is there for anyone who does.", False),
    ("- \"Where it appears\" tells you the screen or lesson, to help with context.", False),
    ("- \"Note from us\" flags rows that need a specific look (pronunciation comparisons, facts "
     "about the author) and rows you can skip (payment and legal text that will be rewritten).", False),
    ("- Rows marked \"Has variants\" change with a number (1 mistake / 2 mistakes). Suggest the "
     "wording; we handle the variants.", False),
    ("- Transliteration (Hebrew written in Latin letters) follows one rule across the course: no "
     "accent marks, \"ch\" for ח and כ (as in Bach). If a spelling looks wrong to you, tell us in Comment.", False),
    ("- The tone is the author's: direct, warm, short sentences, no hype. Natural American English is the goal.", False),
]

COLUNAS = ['ID', 'Where it appears', 'Portuguese (original)', 'English now', 'Note from us',
           'Your version', 'Comment', 'Status']
LARGURAS = [18, 24, 40, 55, 30, 55, 30, 11]
NOTA_PAG = 'Payment or legal text written for Brazil. It will be rewritten for the US launch: you can skip this row.'
NOTA_PRON = 'Pronunciation comparison: does it make sense for an English speaker?'
NOTA_ERRO = ('Common-mistake note: is this a mistake English speakers really make with this letter? '
             'If not, which one do they make?')
NOTA_BIO = 'About the author: improve the English, but do not change the facts.'
BIO = ('venda-trajetoria', 'venda-sobre', 'motor-autor', 'motor-credito')
ROTULO = {'letters': 'Letters', 'modules': 'Modules', 'nikud': 'Vowel signs', 'extras': 'Extra sounds',
          'real-world': 'Real world', 'sentences': 'Sentences', 'eras': 'Eras', 'culture': 'Culture cards',
          'missoes': 'Missions', 'textos': 'Short texts', 'bonus': 'Bonus readings', 'biblia': 'Bible page',
          'imagens-vocabulario': 'Picture descriptions'}


def gerar(pasta, raiz=RAIZ):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    d = dados(raiz)
    telas, conteudo, translit, marcas = d['telas'], d['conteudo'], d['translit'], d['marcas']
    os.makedirs(pasta, exist_ok=True)
    cab, cabfundo = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1D545E')
    amarelo, cinza = PatternFill('solid', fgColor='FFF6CC'), PatternFill('solid', fgColor='F2F2F2')
    quebra, fina = Alignment(wrap_text=True, vertical='top'), Side(style='thin', color='DDDDDD')

    def leia(wb, extra):
        ws = wb.active
        ws.title = 'READ ME'
        ws.column_dimensions['A'].width = 110
        for i, (t, negrito) in enumerate(LEIA + [("", False)] + extra, 1):
            c = ws.cell(i, 1, t)
            c.font = Font(bold=negrito, size=13 if negrito else 11)
            c.alignment = Alignment(wrap_text=True)

    def folha(wb, nome, colunas, larguras, linhas):
        ws = wb.create_sheet(nome[:31])
        ws.append(colunas)
        for j, c in enumerate(ws[1], 1):
            c.font, c.fill, c.alignment = cab, cabfundo, quebra
            ws.column_dimensions[c.column_letter].width = larguras[j - 1]
        dv = DataValidation(type='list', formula1='"OK,Changed,Question"', allow_blank=True)
        ws.add_data_validation(dv)
        iv, ist = colunas.index('Your version') + 1, colunas.index('Status') + 1
        for linha in linhas:
            ws.append(linha)
            r = ws.max_row
            for j in range(1, len(colunas) + 1):
                c = ws.cell(r, j)
                c.alignment, c.border = quebra, Border(bottom=fina)
            ws.cell(r, iv).fill = amarelo
            ws.cell(r, 1).fill = cinza
            ws.cell(r, 1).font = Font(color='888888', size=9)
            dv.add(ws.cell(r, ist))
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        return len(linhas)

    def nota_tela(n, k, v):
        m = marcas.get(f'{n}:{k}', '')
        notas = []
        if m.startswith('etapa 5'):
            notas.append(NOTA_PAG)
        elif m.startswith('etapa 3'):
            notas.append(NOTA_PRON)
        if n in BIO or (n == 'venda-fecho' and k in ('apresentacao', 'tituloProfessor', 'resumo1', 'resumo2', 'resumo3')):
            notas.append(NOTA_BIO)
        if v.get('variantes'):
            notas.append('Has variants (singular/plural).')
        return ' '.join(notas)

    def humano(nome):
        p, _, r = nome.partition('-')
        area = {'venda': 'Sales page', 'conta': 'Account', 'licao': 'Lesson', 'painel': 'Dashboard',
                'academia': 'Academy', 'motor': 'Course', 'sessao': 'Session'}.get(p, p)
        return f"{area} · {r.replace('-', ' ')}" if r else area

    def linhas_tela(nomes):
        out = []
        for n in nomes:
            for k, v in telas[n].items():
                if not re.search(r'[A-Za-z]', v['en'] or ''):
                    continue
                out.append([f'tela:{n}:{k}', humano(n), v['pt'], v['en'], nota_tela(n, k, v), '', '', ''])
        return out

    venda = [n for n in telas if n.split('-')[0] in ('venda', 'geral', 'nao', 'sessao', 'conta')
             or n in ('motor-seo', 'motor-testimonials', 'motor-credito')]
    curso = [n for n in telas if n not in venda]
    total = {}

    wb = Workbook()
    leia(wb, [("This file", True), ("Everything a visitor sees before and around buying: the sales page, the free demo, "
              "sign up, log in, checkout and the account screens. Start here: it is what decides whether someone buys.", False)])
    grupos = {'Sales page': [n for n in venda if not n.startswith(('conta', 'sessao'))],
              'Account and checkout': [n for n in venda if n.startswith(('conta', 'sessao'))]}
    total['1'] = sum(folha(wb, g, COLUNAS, LARGURAS, linhas_tela(ns)) for g, ns in grupos.items())
    wb.save(os.path.join(pasta, '1-Sales-page-and-account.xlsx'))

    wb = Workbook()
    leia(wb, [("This file", True), ("The screens inside the course: lesson steps, buttons, feedback after an answer, "
              "the dashboard, the Academy games and the rewards shop.", False)])
    grupos = {'Lessons': [n for n in curso if n.startswith('licao')], 'Dashboard': [n for n in curso if n.startswith('painel')],
              'Academy': [n for n in curso if n.startswith('academia')], 'Course engine': [n for n in curso if n.startswith('motor')]}
    resto = [n for n in curso if not any(n in v for v in grupos.values())]
    if resto:
        grupos['Other'] = resto
    total['2'] = sum(folha(wb, g, COLUNAS, LARGURAS, linhas_tela(ns)) for g, ns in grupos.items())
    wb.save(os.path.join(pasta, '2-Course-screens.xlsx'))

    def onde(arq, c):
        partes = [p for p in c.split('/') if p]
        return f"{ROTULO[arq]} · {' › '.join(partes[:-1])[:60]}" if len(partes) > 1 else ROTULO[arq]

    def nota_conteudo(a, c, r):
        if a in ('letters', 'extras') and re.search(r'brazilianMistake|paraBrasileiro|/erro/', c):
            return NOTA_ERRO
        if a in ('letters', 'extras') and re.search(r'portugu|french', r['en'], re.I):
            return NOTA_PRON
        return ''

    grupos = {'Letter lessons': ['letters'], 'Modules and signs': ['modules', 'nikud', 'extras'],
              'Reading': ['sentences', 'textos', 'missoes', 'real-world'], 'Culture': ['culture', 'eras'],
              'Bonus readings': ['bonus', 'biblia'], 'Pictures': ['imagens-vocabulario']}
    wb = Workbook()
    leia(wb, [("This file", True), ("The teaching content: how each letter works and sounds, examples, quiz questions and "
              "answers, short reading texts, culture cards, Bible bonus readings and the descriptions of the vocabulary pictures.", False),
              ("Pronunciation notes compare Hebrew sounds with English sounds. If a comparison does not work for your accent, "
               "say so in Comment: this is one of the most useful things you can catch.", False)])
    n3 = 0
    for g, arqs in grupos.items():
        linhas = [[f'{a}:{c}', onde(a, c), r['pt'], r['en'], nota_conteudo(a, c, r), '', '', '']
                  for a in arqs for c, r in conteudo.get(a, {}).items() if r.get('en')]
        n3 += folha(wb, g, COLUNAS, LARGURAS, linhas)
    total['3'] = n3
    wb.save(os.path.join(pasta, '3-Course-content.xlsx'))

    wb = Workbook()
    leia(wb, [("This file", True), ("Every transliteration in the course: the Hebrew word, how it is written in Latin "
              "letters for English speakers, and the Brazilian spelling it came from (for reference). Read the English column "
              "aloud: does it lead an English speaker to the right sound?", False),
              ("Rule used: no accent marks; ch = the throat sound of ח and כ (Bach, not chair); sh = ש; tz for צ; zh for the "
               "ז׳ sound (as in measure); j for ג׳ (as in jam); tch for צ׳ (as in chair).", False)])
    tc = ['ID', 'Hebrew', 'Brazilian spelling (reference)', 'English now', 'Note from us', 'Your version', 'Comment', 'Status']
    linhas = [[f"{t['arquivo']}:{t['caminho']}", t['he'], t['pt'], t['en'], '', '', '', ''] for t in translit]
    total['4'] = folha(wb, 'Transliteration', tc, [18, 22, 26, 26, 10, 28, 30, 11], linhas)
    for r in range(2, len(linhas) + 2):
        c = wb['Transliteration'].cell(r, 2)
        c.font = Font(name='Noto Sans Hebrew', size=14)
        c.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
    wb.save(os.path.join(pasta, '4-Transliteration.xlsx'))
    return total


# ── volta: aplicar o que os revisores escreveram ──────────────────────────

def ler_planilhas(arquivos):
    from openpyxl import load_workbook
    linhas = []
    for arq in arquivos:
        wb = load_workbook(arq, read_only=True, data_only=True)
        for ws in wb.worksheets:
            if ws.title == 'READ ME':
                continue
            it = ws.iter_rows(values_only=True)
            try:
                cab = [str(c or '').strip() for c in next(it)]
            except StopIteration:
                continue
            if 'ID' not in cab or 'Your version' not in cab:
                continue
            col = {n: i for i, n in enumerate(cab)}
            for n, vals in enumerate(it, 2):
                pega = lambda nome: (str(vals[col[nome]]).strip() if nome in col and col[nome] < len(vals)
                                     and vals[col[nome]] is not None else '')
                if not pega('ID'):
                    continue
                linhas.append({'arquivo': os.path.basename(arq), 'folha': ws.title, 'linha': n, 'id': pega('ID'),
                               'agora': pega('English now'), 'versao': pega('Your version'),
                               'comentario': pega('Comment'), 'status': pega('Status')})
    return linhas


def mesmo(re_, a, b):
    return re_.findall(a) == re_.findall(b)


def literal_js(texto):
    """Uma string JS entre aspas simples, como o resto dos dicionários."""
    return "'" + texto.replace('\\', '\\\\').replace("'", "\\'") + "'"


def ler_literal(src):
    """O texto de uma string JS de aspas simples (sem template, sem concatenação)."""
    m = re.fullmatch(r"'((?:[^'\\]|\\.)*)'", src.strip())
    if not m:
        return None
    return re.sub(r"\\(.)", lambda x: {'n': '\n', 't': '\t'}.get(x.group(1), x.group(1)), m.group(1))


def achar_em_tela(src, chave):
    """A linha `chave: '...'` do bloco `en`: (início, fim, texto) ou None."""
    i = src.find('const en')
    if i < 0:
        return None
    padrao = re.compile(r"^(\s+" + re.escape(chave) + r"\s*:\s*)('(?:[^'\\\n]|\\.)*')(\s*,?\s*)$", re.M)
    achados = [m for m in padrao.finditer(src, i)]
    if len(achados) != 1:
        return None
    m = achados[0]
    texto = ler_literal(m.group(2))
    return None if texto is None else (m.start(2), m.end(2), texto)


def aplicar(arquivos, gravar=False, raiz=RAIZ):
    linhas = ler_planilhas(arquivos)
    aplicadas, mao, conflitos, comentarios = [], [], [], []
    conteudo_novo = {}
    telas_novas = {}

    for l in linhas:
        if l['comentario'] or l['status'].lower() == 'question':
            comentarios.append(l)
        versao = l['versao']
        if not versao or versao == l['agora']:
            continue
        lid = l['id']
        if not mesmo(HEBRAICO, l['agora'], versao):
            mao.append((l, 'o hebraico ou as marcas {{...}} mudaram'))
            continue
        if not mesmo(VARIAVEL, l['agora'], versao):
            mao.append((l, 'as variáveis entre chaves mudaram'))
            continue

        if lid.startswith('tela:'):
            _, nome, chave = lid.split(':', 2)
            caminho = next((os.path.join(raiz, 'app/src/i18n/telas', nome + ext) for ext in ('.ts', '.tsx')
                            if os.path.exists(os.path.join(raiz, 'app/src/i18n/telas', nome + ext))), None)
            if not caminho:
                mao.append((l, 'a tela não existe mais'))
                continue
            src = telas_novas.get(caminho) or open(caminho, encoding='utf-8').read()
            achado = achar_em_tela(src, chave)
            if not achado:
                mao.append((l, 'na tela o texto não é uma string simples (função, JSX ou várias linhas)'))
                continue
            ini, fim, hoje = achado
            if hoje == versao:
                continue  # já aplicada numa passada anterior
            if hoje != l['agora']:
                conflitos.append((l, hoje))
                continue
            telas_novas[caminho] = src[:ini] + literal_js(versao) + src[fim:]
            aplicadas.append(l)
            continue

        arq, _, caminho = lid.partition(':')
        if caminho.startswith('/'):
            mao.append((l, 'transliteração: sai da regra de tools/translit-en.mjs'))
            continue
        destino = os.path.join(raiz, 'data/en', arq + '.json')
        if not os.path.exists(destino):
            mao.append((l, 'arquivo desconhecido'))
            continue
        dados_en = conteudo_novo.get(destino)
        if dados_en is None:
            with open(destino, encoding='utf-8') as f:
                dados_en = json.load(f)
        reg = dados_en.get(caminho)
        if reg is None:
            mao.append((l, 'o texto não existe mais no curso'))
            continue
        if reg.get('en') == versao:
            continue  # já aplicada numa passada anterior
        if reg.get('en') != l['agora']:
            conflitos.append((l, reg.get('en') or ''))
            continue
        reg['en'] = versao
        reg.pop('revisar', None)
        if l['comentario']:
            reg['comentario'] = l['comentario']
        conteudo_novo[destino] = dados_en
        aplicadas.append(l)

    if gravar:
        for caminho, src in telas_novas.items():
            with open(caminho, 'w', encoding='utf-8') as f:
                f.write(src)
        for destino, d in conteudo_novo.items():
            with open(destino, 'w', encoding='utf-8') as f:
                f.write(json.dumps(d, ensure_ascii=False, indent=2) + '\n')
    return {'linhas': len(linhas), 'aplicadas': aplicadas, 'mao': mao, 'conflitos': conflitos, 'comentarios': comentarios}


def relatorio(r, gravou):
    onde = lambda l: f"`{l['id']}` ({l['arquivo']}, {l['folha']}, linha {l['linha']})"
    out = ['# Revisão do inglês: o que voltou das planilhas', '',
           f"{r['linhas']} linhas lidas. {len(r['aplicadas'])} versões "
           f"{'aplicadas' if gravou else 'prontas para aplicar (rodar com --gravar)'}, "
           f"{len(r['conflitos'])} em conflito, {len(r['mao'])} para aplicar à mão, "
           f"{len(r['comentarios'])} comentários.", '']
    if r['conflitos']:
        out += ['## O inglês mudou depois de a planilha sair', '',
                'A versão do revisor partiu de um texto que não é mais o de hoje. Decidir à mão.', '']
        for l, hoje in r['conflitos']:
            out += [f"- {onde(l)}", f"  - na planilha: {l['agora']}", f"  - hoje: {hoje}", f"  - revisor: {l['versao']}"]
        out.append('')
    if r['mao']:
        out += ['## Para aplicar à mão', '']
        for l, motivo in r['mao']:
            out += [f"- {onde(l)}: {motivo}", f"  - hoje: {l['agora']}", f"  - revisor: {l['versao']}"]
        out.append('')
    if r['comentarios']:
        out += ['## Comentários e perguntas dos revisores', '']
        for l in r['comentarios']:
            out += [f"- {onde(l)}{' (Question)' if l['status'].lower() == 'question' else ''}",
                    f"  - texto: {l['agora']}", f"  - comentário: {l['comentario'] or '(sem texto)'}"]
        out.append('')
    if r['aplicadas']:
        out += ['## Aplicadas' if gravou else '## Prontas para aplicar', '']
        for l in r['aplicadas']:
            out += [f"- {onde(l)}", f"  - antes: {l['agora']}", f"  - depois: {l['versao']}"]
        out.append('')
    return '\n'.join(out)


# ── a prova ──────────────────────────────────────────────────────────────

def ensaiar():
    """Gera as planilhas numa cópia do repositório, preenche como um revisor
    faria e aplica: cada caso cai onde devia, e nada fora da cópia muda."""
    from openpyxl import load_workbook
    tmp = tempfile.mkdtemp()
    copia = os.path.join(tmp, 'repo')
    ignorar = shutil.ignore_patterns('node_modules', '.next', 'out', 'out-en', '.git', '.wrangler', 'audio', 'public',
                                     'editorial', 'raizes', 'content')
    shutil.copytree(RAIZ, copia, ignore=ignorar)
    os.symlink(os.path.join(RAIZ, 'app', 'node_modules'), os.path.join(copia, 'app', 'node_modules'))
    pasta = os.path.join(tmp, 'planilhas')
    gerar(pasta, copia)

    casos = {}

    def preencher(nome, alvo, versao, comentario=''):
        caminho = os.path.join(pasta, nome)
        wb = load_workbook(caminho)
        for ws in wb.worksheets[1:]:
            cab = [c.value for c in ws[1]]
            iid, iagora, iv, ic = (cab.index(x) + 1 for x in ('ID', 'English now', 'Your version', 'Comment'))
            for r in range(2, ws.max_row + 1):
                lid, agora = ws.cell(r, iid).value or '', ws.cell(r, iagora).value or ''
                if alvo(lid, agora):
                    v = versao(agora) if callable(versao) else versao
                    ws.cell(r, iv).value = v
                    ws.cell(r, ic).value = comentario or None
                    wb.save(caminho)
                    return lid, agora, v
        raise SystemExit(f'ensaio: nenhuma linha para o caso em {nome}')

    # uma string simples de tela
    casos['tela'] = preencher('2-Course-screens.xlsx',
                              lambda i, a: i.startswith('tela:') and re.fullmatch(r"[A-Za-z ,.]{8,40}", a or '') is not None,
                              lambda a: a + ' (reviewed)')
    # uma função de tela (vai para a mão)
    casos['funcao'] = preencher('2-Course-screens.xlsx', lambda i, a: i.startswith('tela:') and '{' in (a or '')
                                and '{{' not in a, lambda a: a.replace(' ', '  ', 1))
    # um texto do curso, com comentário
    casos['conteudo'] = preencher('3-Course-content.xlsx',
                                  lambda i, a: not i.startswith('tela:') and '{{' not in (a or '') and len(a or '') > 20,
                                  lambda a: 'Reviewed: ' + a, 'Sounds more natural.')
    # um texto do curso que mexe no hebraico (recusado)
    casos['hebraico'] = preencher('3-Course-content.xlsx', lambda i, a: '{{' in (a or ''),
                                  lambda a: re.sub(r'\{\{[^}]*\}\}', '', a, count=1))
    # uma transliteração (vai para a mão)
    casos['translit'] = preencher('4-Transliteration.xlsx', lambda i, a: bool(a), lambda a: a + 'h')

    arqs = [os.path.join(pasta, f) for f in sorted(os.listdir(pasta))]
    antes = {f: open(os.path.join(RAIZ, f), 'rb').read() for f in ('data/en/letters.json',)}
    r = aplicar(arqs, gravar=True, raiz=copia)
    ids = lambda xs: {x[0]['id'] if isinstance(x, tuple) else x['id'] for x in xs}
    erros = []
    if casos['tela'][0] not in ids(r['aplicadas']):
        erros.append('a string de tela não foi aplicada')
    if casos['conteudo'][0] not in ids(r['aplicadas']):
        erros.append('o texto do curso não foi aplicado')
    for k in ('funcao', 'hebraico', 'translit'):
        if casos[k][0] not in ids(r['mao']):
            erros.append(f'o caso {k} não foi para a lista do que aplicar à mão')
    if casos['conteudo'][0] not in ids(r['comentarios']):
        erros.append('o comentário do revisor não entrou no relatório')
    # o que foi gravado está mesmo na cópia
    _, nome, chave = casos['tela'][0].split(':', 2)
    src = ''.join(open(os.path.join(copia, 'app/src/i18n/telas', nome + e), encoding='utf-8').read()
                  for e in ('.ts', '.tsx') if os.path.exists(os.path.join(copia, 'app/src/i18n/telas', nome + e)))
    if literal_js(casos['tela'][2]) not in src:
        erros.append('a tela gravada não traz a versão do revisor')
    arq, _, cam = casos['conteudo'][0].partition(':')
    reg = json.load(open(os.path.join(copia, 'data/en', arq + '.json'), encoding='utf-8'))[cam]
    if reg['en'] != casos['conteudo'][2] or reg.get('comentario') != 'Sounds more natural.':
        erros.append('o data/en gravado não traz a versão e o comentário')
    # aplicar de novo a mesma planilha não muda nada (o "English now" já é outro)
    r2 = aplicar(arqs, gravar=False, raiz=copia)
    if r2['aplicadas'] or r2['conflitos']:
        erros.append('aplicar duas vezes não reconheceu o que já estava aplicado')
    if any(open(os.path.join(RAIZ, f), 'rb').read() != b for f, b in antes.items()):
        erros.append('o repositório de verdade mudou')
    shutil.rmtree(tmp)
    if erros:
        print('ensaio falhou:\n  - ' + '\n  - '.join(erros))
        sys.exit(1)
    print(f"ensaio: {len(r['aplicadas'])} aplicadas, {len(r['mao'])} à mão, {len(r['comentarios'])} comentário; "
          "aplicar de novo não muda nada.")


def principal(args):
    if not args:
        print(__doc__)
        sys.exit(1)
    cmd, resto = args[0], args[1:]
    if cmd == 'gerar' and resto:
        print(gerar(resto[0]))
    elif cmd == 'aplicar' and resto:
        gravar = '--gravar' in resto
        saida = resto[resto.index('--relatorio') + 1] if '--relatorio' in resto else None
        arqs = [a for i, a in enumerate(resto) if a.endswith('.xlsx')]
        r = aplicar(arqs, gravar=gravar)
        texto = relatorio(r, gravar)
        if saida:
            with open(saida, 'w', encoding='utf-8') as f:
                f.write(texto + '\n')
        print(texto)
    elif cmd == 'ensaiar':
        ensaiar()
    else:
        print(__doc__)
        sys.exit(1)


if __name__ == '__main__':
    principal(sys.argv[1:])
